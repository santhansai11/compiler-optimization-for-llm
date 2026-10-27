import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np

from utils.benchmarking import paired_summary, randomized_benchmark, suite_summary, write_json


class TimingTests(unittest.TestCase):
    def test_randomized_rounds_keep_pairing_and_units(self):
        counts = {"a": 0, "b": 0, "c": 0}

        def call(name):
            counts[name] += 1
            return name

        calls = {k: lambda k=k: call(k) for k in counts}
        with patch("utils.benchmarking.time.perf_counter_ns", side_effect=range(0, 180000, 1000)):
            record, outputs = randomized_benchmark(calls, rounds=30, warmup=2, seed=5)
        self.assertEqual(counts, dict.fromkeys(counts, 32))
        self.assertEqual(outputs, dict(zip(counts, counts)))
        self.assertEqual(record["unit"], "us")
        self.assertEqual(record["samples_us"], {k: [1.0]*30 for k in counts})
        self.assertGreater(len({tuple(x) for x in record["round_order"]}), 1)
        for order in record["round_order"]:
            self.assertEqual(set(order), set(counts))
        with self.assertRaises(ValueError):
            randomized_benchmark(calls, rounds=29)

    def test_paired_interval_direction_and_exact_zero_differences(self):
        fast = paired_summary([100]*30, [80]*30)
        self.assertEqual(fast["median_paired_difference_us"], 20)
        self.assertEqual(fast["speedup_ratio_of_medians"], 1.25)
        self.assertEqual(fast["paired_difference_ci95_us"], [20, 20])
        self.assertEqual(fast["ci_direction"], "faster")
        self.assertFalse(fast["observed_slowdown"])
        same = paired_summary([100]*30, [100]*30)
        self.assertEqual(same["wilcoxon"]["pvalue_two_sided"], 1)
        self.assertEqual(same["ci_direction"], "inconclusive")
        slow = paired_summary([80]*30, [100]*30)
        self.assertTrue(slow["observed_slowdown"])
        self.assertEqual(slow["ci_direction"], "slower")

    def test_invalid_samples_and_nonfinite_json_rejected(self):
        for a,b in [([],[]),([1],[1,2]),([0],[1]),([float("nan")],[1])]:
            with self.assertRaises(ValueError):
                paired_summary(a,b)
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaises(ValueError):
                write_json(Path(folder)/"bad.json", {"value":float("nan")})

    def test_suite_counts_workloads_not_repetitions(self):
        rows = [{"timing":{"median_us":{"original":a,"guided":b}}} for a,b in [(10,5),(10,20),(10,10)]]
        result = suite_summary(rows,"original","guided")
        self.assertEqual(result["pairs"], 3)
        self.assertEqual(result["observed_slowdowns"], 1)
        self.assertEqual(result["exact_median_ties"], 1)
        self.assertAlmostEqual(result["geomean_speedup"], 1)


class XlaEvidenceTests(unittest.TestCase):
    def test_hlo_comparison_ignores_debug_but_preserves_operations(self):
        from utils.xla_runtime import normalized_hlo
        left = 'HloModule a, is_scheduled=true\n\nFileNames\n1 "first.py"\n\nStackFrames\n1 {file_location_id=1}\n\nENTRY main {\n ROOT x = f32[] add(a,b), metadata={stack_frame_id=1}\n}\n'
        right = left.replace('HloModule a,','HloModule b,').replace('first.py','second.py').replace('file_location_id=1','file_location_id=2')
        self.assertEqual(normalized_hlo(left),normalized_hlo(right))
        self.assertNotEqual(normalized_hlo(left),normalized_hlo(right.replace('add(a,b)','multiply(a,b)')))

    def test_each_lab_and_rewritten_graph_matches_numpy(self):
        from models.operations import OPERATION_DEMOS
        from pipeline import run_pipeline
        from utils.xla_runtime import XlaProgram
        for name, (_, build, _) in OPERATION_DEMOS.items():
            graph = build(1, 8, 16)
            original = XlaProgram(graph)
            compiled = XlaProgram(run_pipeline(graph, partition=False)[0])
            np.testing.assert_allclose(original.output, compiled.output, rtol=1e-4, atol=1e-5, err_msg=name)
            self.assertIn("HloModule", original.hlo)
            self.assertIn("temp_size_in_bytes", original.record()["memory"])

    def test_checkpoint_features_do_not_include_gpu_cost_priors(self):
        from benchmarks.xla_guidance_study import features
        from models.operations import attention_demo
        from search import learned_optimizer as lo
        x, _, glob = features(attention_demo())
        self.assertTrue(np.all(x[:,len(lo.OPS)+2] == 0))
        self.assertTrue(np.all(glob[5:8] == 0))
        self.assertEqual(glob[-1], 0)

    def test_decision_quality_counts_missing_predictions_as_rejections(self):
        from benchmarks.xla_guidance_study import decision_metrics
        rows = [
            {"predicted_speedup":1.2,"measured_speedup":1.2,"predicted_accept":True},
            {"predicted_speedup":1.2,"measured_speedup":0.9,"predicted_accept":True},
            {"predicted_speedup":None,"measured_speedup":1.2,"predicted_accept":False},
            {"predicted_speedup":1.0,"measured_speedup":1.0,"predicted_accept":False}]
        result = decision_metrics(rows,1.05)
        for key in ("correct_acceptances","wrong_acceptances","missed_beneficial_rewrites","correct_rejections","missing_estimates"):
            self.assertEqual(result[key],1)
        self.assertEqual(result["gate_accuracy"],0.5)
        self.assertEqual(result["missing_estimate_rejection_rate"],0.25)

    def test_invalid_estimates_never_authorize_rewrites(self):
        from benchmarks.xla_guidance_study import guided
        from models.operations import scale_chain_demo

        class MissingModel:
            def predict_log_us(self, feats):
                return np.full(len(feats), np.nan)

        original = scale_chain_demo()
        result, trace, decisions, _ = guided(original,MissingModel(),4,12,1.05)
        self.assertEqual(result.node_count(),original.node_count())
        self.assertEqual(trace,[])
        self.assertTrue(decisions)
        self.assertTrue(all(d["rejection_reason"] == "missing_estimate" for d in decisions))

    def test_dashboard_export_preserves_samples_environment_and_hlo(self):
        from xla_compare import compare_attention_xla
        from utils.reports import write_report
        from openpyxl import load_workbook
        result = compare_attention_xla(batch=1,seq=8,d_model=16,warmup=1,runs=3)
        self.assertEqual(result["timed_runs"],30)
        self.assertTrue(result["all_correct"])
        self.assertIn("cpu",result["backends"][-1]["device"].lower())
        self.assertIn("HloModule",result["xla_optimized_hlo"])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"evidence.xlsx"
            write_report(path,{"xla":result})
            wb = load_workbook(path,read_only=True)
            self.assertEqual(wb["XLA raw timings"].max_row,91)
            self.assertIn("XLA paired statistics",wb.sheetnames)
            self.assertIn("XLA environment",wb.sheetnames)
            self.assertIn("XLA optimized HLO",wb.sheetnames)
            wb.close()


if __name__ == "__main__":
    unittest.main()
