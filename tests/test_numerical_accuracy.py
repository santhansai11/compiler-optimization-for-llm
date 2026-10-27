import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np

from ir.graph import ComputationGraph
from pipeline import run_pipeline
from utils.executor import _apply_steps, _make_feed, _silu, execute
from utils.metrics import compute_metrics
from utils.validation import compare_outputs


def graph_with_input():
    graph = ComputationGraph('regression')
    graph.meta['config'] = {'batch': 1, 'seq': 2, 'd_model': 2}
    graph.add_operation('x', 'Input')
    return graph


class NumericalMetricTests(unittest.TestCase):
    def test_element_percentage_mae_and_rmse_have_real_definitions(self):
        result = compare_outputs({'y': np.array([1, 2, 3, 4])},
                                 {'y': np.array([1, 2, 4, 4])}, rtol=0, atol=0)
        self.assertFalse(result['preserved'])
        self.assertEqual(result['preservation_pct'], 75)
        self.assertEqual(result['mae'], .25)
        self.assertEqual(result['rmse'], .5)

    def test_every_output_is_checked(self):
        result = compare_outputs({'a': np.ones(2), 'b': np.zeros(2)},
                                 {'a': np.ones(2), 'b': np.ones(2)})
        self.assertFalse(result['preserved'])
        self.assertEqual(result['preservation_pct'], 50)

    def test_nonfinite_missing_and_shape_mismatch_are_json_safe_failures(self):
        for candidate in ({'y': np.array([np.nan])}, {'y': np.array([np.inf])},
                          {'y': np.array([1., 2.])}, {}):
            result = compare_outputs({'y': np.array([1.])}, candidate)
            self.assertFalse(result['preserved'])
            self.assertIsNone(result['max_abs_diff'])
            json.dumps(result, allow_nan=False)
        self.assertFalse(compare_outputs({'y': np.array([np.inf])},
                                         {'y': np.array([np.inf])})['preserved'])

    def test_tolerance_uses_original_as_reference(self):
        self.assertFalse(compare_outputs({'y': np.array([1.])},
                                         {'y': np.array([2.])}, rtol=.75, atol=0)['preserved'])

    def test_aliases_count_unique_storage_and_intermediate_outputs_survive(self):
        g = graph_with_input()
        g.graph.nodes['x']['is_output'] = True
        g.add_operation('view', 'Reshape', shape_arg=(1, 4), is_output=True)
        g.add_dependency('x', 'view')
        feed = np.arange(4, dtype=np.float32).reshape(1, 2, 2)
        result = execute(g, repeats=1, warmup=0, feed=feed)
        self.assertEqual(set(result['outputs']), {'x', 'view'})
        self.assertEqual(result['peak_bytes'], feed.nbytes)

    def test_latency_is_median_not_fastest_sample(self):
        g = graph_with_input()
        with patch('utils.executor.time.perf_counter', side_effect=[0, 10, 20, 22, 30, 34]):
            result = execute(g, repeats=3, warmup=0)
        self.assertEqual(result['latency_s'], 4)
        self.assertEqual(result['timings_s'], [10, 2, 4])

    def test_integer_tokens_are_not_rounded_through_float32(self):
        g = graph_with_input()
        g.meta['config']['llm'] = True
        ids = np.array([[16777217]], dtype=np.int64)
        np.testing.assert_array_equal(_make_feed(g, ids), ids)

    def test_stable_silu_and_recorded_projection_bias(self):
        x = np.array([-1000., -1., 0., 1., 1000.], dtype=np.float32)
        with np.errstate(over='raise', invalid='raise'):
            out = _silu(x)
        self.assertTrue(np.isfinite(out).all())
        self.assertAlmostEqual(float(out[-1]), 1000.)
        w = np.eye(2, dtype=np.float32)
        b = np.array([2., -3.], dtype=np.float32)
        np.testing.assert_array_equal(_apply_steps(w, [('mm', w, b)]), w + b)

    def test_throughput_uses_actual_batch_and_exports_raw_evidence(self):
        from openpyxl import load_workbook
        from utils.reports import write_report
        g = graph_with_input()
        g.meta['config']['batch'] = 7
        metrics = compute_metrics(g, g, {}, 0., batch_size=99)
        self.assertEqual(metrics['throughput']['batch'], 7)
        self.assertAlmostEqual(metrics['throughput']['original'] * metrics['latency_ms']['original'] / 1000, 7)
        self.assertEqual(metrics['counts']['kernel_launches_original'], 0)
        self.assertEqual(metrics['modeled']['memory_reduction_pct'], 0)
        self.assertEqual(metrics['timing']['rounds'], 40)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'metrics.xlsx'
            write_report(path, {'compiler_metrics': metrics})
            book = load_workbook(path, read_only=True)
            self.assertEqual(book['CPU raw timings'].max_row, 81)
            self.assertIn('Output validation', book.sheetnames)
            book.close()

    def test_fx_sweep_uses_paired_timings_and_actual_minimum_rounds(self):
        from run_attention_optimization import run_case
        result = run_case(1, 8, 16, warmup=1, runs=3)
        self.assertTrue(result['correct'])
        self.assertEqual(result['timing']['rounds'], 30)
        self.assertEqual(result['paired_comparison']['pairs'], 30)
        self.assertAlmostEqual(result['original_latency'] * 1000,
                               result['timing']['median_us']['original'])

    def test_modeled_memory_percent_uses_same_live_range_policy(self):
        g = graph_with_input()
        g.add_operation('scale', 'Scale', factor=1., is_output=True)
        g.add_dependency('x', 'scale')
        metrics = compute_metrics(g, graph_with_input(), {}, 0.)
        self.assertEqual(metrics['modeled']['memory_original_mb'], 72.)
        self.assertEqual(metrics['modeled']['memory_optimized_mb'], 48.)
        self.assertAlmostEqual(metrics['modeled']['memory_reduction_pct'], 100. / 3.)


class RewriteRegressionTests(unittest.TestCase):
    def assert_equivalent(self, graph, **flags):
        optimized, _, _ = run_pipeline(graph, **flags)
        before = execute(graph, repeats=1, warmup=0)['outputs']
        after = execute(optimized, repeats=1, warmup=0)['outputs']
        result = compare_outputs(before, after)
        self.assertTrue(result['preserved'], result)
        return optimized

    def test_identity_rewrite_preserves_subtraction_operand_order(self):
        g = graph_with_input()
        g.add_operation('id', 'Identity')
        g.add_operation('c', 'Constant', value=np.float32(2))
        g.add_operation('out', 'Sub', is_output=True)
        for a, b in [('x', 'id'), ('id', 'out'), ('c', 'out')]:
            g.add_dependency(a, b)
        self.assert_equivalent(g)

    def test_identity_cannot_collapse_two_operands_or_remove_output(self):
        g = graph_with_input()
        g.add_operation('id', 'Identity', is_output=True)
        g.add_operation('sum', 'Add', is_output=True)
        for a, b in [('x', 'id'), ('x', 'sum'), ('id', 'sum')]:
            g.add_dependency(a, b)
        self.assert_equivalent(g)
        g.graph.nodes['id'].pop('is_output')
        self.assert_equivalent(g)

    def test_cse_preserves_operand_order(self):
        g = graph_with_input()
        for name in ['a', 'b']:
            g.add_operation(name, 'Scale', factor=2.)
            g.add_dependency('x', name)
        g.add_operation('left', 'Relu', is_output=True)
        g.add_operation('c', 'Constant', value=np.float32(3))
        g.add_operation('right', 'Sub', is_output=True)
        for a, b in [('a', 'left'), ('b', 'right'), ('c', 'right')]:
            g.add_dependency(a, b)
        self.assert_equivalent(g)

    def test_cse_does_not_merge_different_exponents_or_declared_outputs(self):
        g = graph_with_input()
        for name, exponent in [('a', 2), ('b', 3)]:
            g.add_operation(name, 'Pow', exponent=exponent)
            g.add_dependency('x', name)
        g.add_operation('out', 'Sub', is_output=True)
        g.add_dependency('a', 'out')
        g.add_dependency('b', 'out')
        self.assert_equivalent(g)
        g = graph_with_input()
        for name in ['a', 'b']:
            g.add_operation(name, 'Scale', factor=2., is_output=True)
            g.add_dependency('x', name)
        self.assert_equivalent(g)

    def test_two_input_matmul_is_not_a_single_input_linear_projection(self):
        g = graph_with_input()
        g.add_operation('w', 'Constant', value=np.array([[1., 2.], [3., 4.]], dtype=np.float32))
        g.add_operation('mm', 'MatMul')
        g.add_operation('bias', 'Constant', value=np.float32(2))
        g.add_operation('out', 'Add', is_output=True)
        for a, b in [('x', 'mm'), ('w', 'mm'), ('mm', 'out'), ('bias', 'out')]:
            g.add_dependency(a, b)
        self.assert_equivalent(g, normalize=False, canonicalize=False)
        from search.rewrites import enumerate_rewrites
        self.assertFalse(any(rule.name == 'fuse-pair' for rule, _ in enumerate_rewrites(g)))

    def test_fusion_preserves_declared_producer_output(self):
        g = graph_with_input()
        g.add_operation('mm', 'MatMul', is_output=True)
        g.add_operation('out', 'GELU', is_output=True)
        g.add_dependency('x', 'mm')
        g.add_dependency('mm', 'out')
        self.assert_equivalent(g)

    def test_default_constant_bias_and_fusion_node_count(self):
        g = graph_with_input()
        g.add_operation('mm', 'MatMul')
        g.add_operation('bias', 'Constant')
        g.add_operation('out', 'Add', is_output=True)
        for a, b in [('x', 'mm'), ('mm', 'out'), ('bias', 'out')]:
            g.add_dependency(a, b)
        optimized, infos, _ = run_pipeline(g, normalize=False, canonicalize=False)
        self.assertTrue(compare_outputs(execute(g, repeats=1)['outputs'],
                                        execute(optimized, repeats=1)['outputs'])['preserved'])
        self.assertEqual(infos['semantic_merging']['nodes_removed'], 2)

    def test_constant_fold_honours_transpose_rhs(self):
        g = ComputationGraph()
        g.add_operation('a', 'Constant', value=np.array([[1., 2.], [3., 4.]], dtype=np.float32))
        g.add_operation('b', 'Constant', value=np.array([[2., 0.], [1., 3.]], dtype=np.float32))
        g.add_operation('out', 'MatMul', transpose_rhs=True, is_output=True)
        g.add_dependency('a', 'out')
        g.add_dependency('b', 'out')
        self.assert_equivalent(g)

    def test_normalized_names_remain_unique_and_preserve_seeded_weights(self):
        g = graph_with_input()
        for name in ['a-b', 'a_b']:
            g.add_operation(name, 'MatMul', is_output=True)
            g.add_dependency('x', name)
        self.assert_equivalent(g)

    def test_attention_fusion_retains_output_and_external_operand_order(self):
        from models.operations import attention_demo
        g = attention_demo(batch=1, seq=4, d_model=4)
        output = next(n for n, d in g.nodes(data=True) if d.get('is_output'))
        g.add_operation('tail', 'Sub', is_output=True)
        g.add_dependency(output, 'tail')
        g.add_dependency('Input', 'tail')
        self.assert_equivalent(g)


if __name__ == '__main__':
    unittest.main()
