"""Excel report writer (openpyxl) for dashboard runs.

One workbook per completed comparison/sweep. Every metric row states whether
it is *measured* (NumPy reference executor / FX / XLA timings), *modeled*
(hand-written GPU cost table) or *structural* (graph counts), so the
workbook never mixes provenance silently.
"""

import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="3B2FA8")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, head in enumerate(header, start=1):
        longest = max([len(str(head))] + [len(str(r[i - 1])) for r in rows
                                         if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(
            60, (widths or {}).get(i, longest + 3))
    ws.freeze_panes = "A2"
    return ws


def metric_rows(m):
    """The ten required output metrics as (name, original, optimized, change, kind)."""
    c, lat, thr = m["counts"], m["latency_ms"], m["throughput"]
    mem, mod, acc = m["peak_memory_mb"], m["modeled"], m["accuracy"]
    acr = m.get("attention_canonicalization_rate")
    error = acc["max_abs_diff"]
    error_text = f"max |diff| {error:.2e}" if error is not None else "Invalid output; error unavailable"
    return [
        ["Inference latency (ms)", lat["original"], lat["optimized"],
         f"{m['speedup']:.3f}x", "measured (NumPy executor)"],
        ["Inference latency, calibrated model (ms)", mod["latency_original_us"] / 1e3,
         mod["latency_optimized_us"] / 1e3, f"{mod['speedup']:.3f}x", "modeled"],
        ["Throughput (samples/s)", thr["original"], thr["optimized"],
         f"{thr['optimized'] / thr['original']:.3f}x", "measured"],
        ["Speedup (measured)", 1.0, m["speedup"], f"{m['speedup']:.3f}x", "measured"],
        ["Speedup (modeled, GPU-calibrated)", 1.0, mod["speedup"], f"{mod['speedup']:.3f}x", "modeled"],
        ["Live graph storage (MiB)", mem["original"], mem["optimized"],
         f"{(1 - mem['optimized'] / mem['original']) * 100:.1f}% lower" if mem["original"] else "",
         "unique backing storage; excludes kernel temporaries and allocator overhead"],
        ["Graph Reduction Ratio (GRR)", "", m["graph_reduction_ratio"],
         f"{c['nodes_original']} -> {c['nodes_optimized']} nodes", "structural"],
        ["Output elements within tolerance (%)", 100.0, acc["preservation_pct"],
         error_text, "numerical equivalence; not dataset accuracy"],
        ["Operator Merge Ratio (OMR)", "", m["operator_merge_ratio"], "", "structural"],
        ["Attention Canonicalization Rate (ACR)", "",
         acr if acr is not None else "N/A",
         f"{m['attention_subgraphs']['canonicalized']}/{m['attention_subgraphs']['detected']} patterns", "structural"],
        ["Kernel launch reduction", c["kernel_launches_original"], c["kernel_launches_optimized"],
         f"{m['kernel_launch_reduction'] * 100:.1f}%", "structural proxy"],
        ["Compilation time (ms)", "", m["compilation_time_s"] * 1e3, "", "measured"],
    ]


def write_report(path, report):
    wb = Workbook()
    wb.remove(wb.active)

    meta = [
        ["Run ID", report.get("run_id", "")],
        ["Created (UTC)", report.get("created_utc", "")],
        ["Report kind", report.get("kind", "")],
        ["Model", report.get("model_name") or report.get("model_id", "")],
        ["Workload", str(report.get("workload", ""))],
        ["Device", report.get("device", "CPU")],
        ["Passes", ", ".join(k for k, v in (report.get("passes") or {}).items() if v)],
    ]
    _sheet(wb, "Run", ["Field", "Value"], meta)

    if report.get("compiler_metrics"):
        metrics = report["compiler_metrics"]
        _sheet(wb, "Metrics", ["Metric", "Original", "Optimized", "Change", "Kind"],
               metric_rows(metrics))
        timing = metrics.get("timing")
        if timing:
            _sheet(wb, "CPU raw timings", ["Round", "Execution position", "Variant", "Latency us"],
                   [[i + 1, pos + 1, name, timing["samples_us"][name][i]]
                    for i, order in enumerate(timing["round_order"]) for pos, name in enumerate(order)])
            _sheet(wb, "CPU timing protocol", ["Field", "Value"],
                   [[k, json.dumps(v)] for k, v in timing.items() if k not in {"round_order", "samples_us"}])
        _sheet(wb, "CPU paired statistics", ["Statistic", "Value"],
               [[k, json.dumps(v)] for k, v in metrics.get("paired_comparison", {}).items()])
        _sheet(wb, "Output validation", ["Statistic", "Value"],
               [[k, json.dumps(v)] for k, v in metrics["accuracy"].items()])
    if report.get("stages"):
        _sheet(wb, "Pipeline stages",
               ["Stage", "Nodes", "Kernel launches", "Modeled latency (us)"],
               [[s["stage"], s["ops"], s.get("launches", ""),
                 s.get("modeled_latency_us", "")] for s in report["stages"]])
    if report.get("architectures"):
        _sheet(wb, "Architectures", ["Architecture", "Nodes before", "Nodes after",
                                     "GRR %", "OMR %", "ACR %", "Kernel launch red. %",
                                     "Latency orig ms", "Latency opt ms", "Speedup (meas.)",
                                     "Speedup (model)", "Throughput orig", "Throughput opt",
                                     "Mem orig MB", "Mem opt MB", "GPU mem orig MB",
                                     "GPU mem opt MB", "Accuracy ok", "Max |diff|",
                                     "Compile ms"],
               [[a["name"], a["nodes_original"], a["nodes_optimized"], a["grr"], a["omr"],
                 a["acr"], a["klr"], a["latency_original_ms"], a["latency_optimized_ms"],
                 a["speedup"], a["modeled_speedup"], a["throughput_original"],
                 a["throughput_optimized"], a["memory_original_mb"],
                 a["memory_optimized_mb"], a["gpu_memory_original_mb"],
                 a["gpu_memory_optimized_mb"], a["accuracy_preserved"], a["max_abs_diff"],
                 a["compile_ms"]] for a in report["architectures"]])
    xla = report.get("xla")
    if xla:
        rows = [[b["name"], b["latency_ms"], b["throughput_tokens_s"],
                 b.get("compile_ms", ""), b["max_abs_error_vs_torch"],
                 b["correct_vs_torch"], b.get("stablehlo_ops", ""), b.get("device", "")]
                for b in xla["backends"]]
        _sheet(wb, "XLA", ["Backend", "Latency ms", "Tokens/s", "Compile ms",
                           "Max err vs torch", "Correct", "StableHLO ops", "Device"], rows)
        timing = xla.get("timing", {})
        if timing:
            _sheet(wb, "XLA raw timings", ["Round", "Execution position", "Backend", "Latency us"],
                   [[i + 1, pos + 1, name, timing["samples_us"][name][i]]
                    for i, order in enumerate(timing["round_order"]) for pos, name in enumerate(order)])
            _sheet(wb, "XLA timing protocol", ["Field", "Value"],
                   [[k, json.dumps(v)] for k, v in timing.items() if k not in {"round_order", "samples_us"}])
        if xla.get("paired_comparisons"):
            _sheet(wb, "XLA paired statistics", ["Comparison", "Statistic", "Value"],
                   [[name, key, json.dumps(value)] for name, stats in xla["paired_comparisons"].items()
                    for key, value in stats.items()])
        if xla.get("environment"):
            _sheet(wb, "XLA environment", ["Field", "Value"],
                   [[k, json.dumps(v)] for k, v in xla["environment"].items()])
        if xla.get("xla_optimized_hlo"):
            _sheet(wb, "XLA optimized HLO", ["Line", "Text"],
                   [[i+1, line] for i,line in enumerate(xla["xla_optimized_hlo"].splitlines())])
    sweep = (report.get("sweep") or {}).get("results")
    if sweep:
        keys = list(sweep[0].keys())
        _sheet(wb, "Workload sweep", keys,
               [[json.dumps(r[k]) if isinstance(r[k], (dict, list)) else r[k] for k in keys] for r in sweep])
    wb.save(path)
